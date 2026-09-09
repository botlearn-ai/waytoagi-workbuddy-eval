"""用合成办公文件验证文本完整性和读取边界。"""

import zipfile

import openpyxl
import pytest
from docx import Document

import extract as readers


def test_docx_reads_body_despite_malformed_unused_relationship(tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("A complete paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Label"
    table.cell(0, 1).text = "42"
    doc.save(path)
    with zipfile.ZipFile(path) as z:
        contents = {n: z.read(n) for n in z.namelist()}
    contents["word/_rels/document.xml.rels"] = b'<Relationships broken="yes"oops>'
    with zipfile.ZipFile(path, "w") as z:
        for name, data in contents.items():
            z.writestr(name, data)
    assert readers.extract(path) == "A complete paragraph\n[表格] Label | 42"


def test_malformed_docx_body_is_reported(tmp_path):
    path = tmp_path / "bad.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", b"<bad")
    with pytest.raises(readers.ExtractError):
        readers.extract(path)


def test_reference_extraction_can_use_larger_explicit_limit(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["start", "x" * 500])
    wb.active.append(["end-marker"])
    wb.save(path)
    assert "end-marker" in readers.extract(path, max_chars=1000)
    assert "end-marker" not in readers.extract(path, max_chars=100)


def test_cell_limit_reports_failure_instead_of_partial_evidence(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append([1, 2, 3])
    wb.save(path)
    monkeypatch.setattr(readers, "_MAX_CELLS", 2)
    with pytest.raises(readers.ExtractError):
        readers.extract(path)


def test_docx_preserves_tabs_and_line_breaks(tmp_path):
    path = tmp_path / "formatted.docx"
    doc = Document()
    doc.add_paragraph("left\tright\nend")
    doc.save(path)
    assert readers.extract(path) == "left\tright\nend"


def test_large_spreadsheet_keeps_last_row_for_comparison(tmp_path):
    path = tmp_path / "large.xlsx"
    wb = openpyxl.Workbook()
    for _ in range(25):
        wb.active.append(["x" * 20000])
    wb.active.append(["END-SOURCE"])
    wb.save(path)
    assert readers.extract(path).endswith("END-SOURCE")
