"""Deliverable degradation tool for the discrimination-check gauge (spec U7).

`make_degraded` takes a gold deliverable and produces three variants used to
calibrate the judging pipeline (U8): a round-trip control (measures a
format library's own load/save loss, no semantic change), a truncated
variant (half the content removed) and a shuffled-numbers variant (numeric
content changed while structure is preserved). This module is itself a
measurement instrument for the calibration run's discriminative validity,
so every non-control variant self-checks that it actually differs from the
gold file before returning — a variant that is accidentally equivalent to
the original would silently defeat the whole calibration gate, so it is
never returned; `DegradeError` is raised instead.

Secrecy: this module only ever touches a deliverable file's own bytes
(never criterion text, prompts, task_ids or rubric_item_ids), and
`DegradeError` messages carry only the variant name, format name and
plain integers/counts — never file content.
"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import openpyxl
import pypdf
from docx import Document as DocxDocument
from docx.oxml.ns import qn as docx_qn
from pptx import Presentation
from pptx.oxml.ns import qn as pptx_qn

from gdpval_eval.extract import extract_document
from gdpval_eval.models import GdpvalEvalError

_SUPPORTED_FORMATS = ("xlsx", "docx", "pptx", "pdf")
_NUMBER_TOKEN_RE = re.compile(r"\d+(?:\.\d+)?")


class DegradeError(GdpvalEvalError):
    """A degraded variant could not be produced, or would be equivalent to
    the gold file. Messages carry only the variant name, format name and
    plain numbers — never file content.
    """


# --------------------------------------------------------------------------
# Shared halving rule for the "truncated" variant (and the pdf fallback the
# "shuffled_numbers" variant uses, since pdf content streams can't be edited
# in place). Drops the back `n // 2` items but always keeps at least one.
# --------------------------------------------------------------------------


def _half_drop_keep(n: int) -> tuple[int, int]:
    if n <= 0:
        return 0, 0
    drop = min(n // 2, n - 1)
    return drop, n - drop


def _assert_count_reduced(before: int, after: int, fmt: str, unit: str) -> None:
    if not after < before:
        raise DegradeError(
            f"truncated: {fmt} {unit} count did not decrease ({before} -> {after})"
        )


def _number_multiset(text: str) -> Counter[str]:
    return Counter(_NUMBER_TOKEN_RE.findall(text))


def _assert_numbers_changed(gold: Path, out_path: Path, fmt: str) -> None:
    gold_numbers = _number_multiset(extract_document(gold).text)
    out_numbers = _number_multiset(extract_document(out_path).text)
    if gold_numbers == out_numbers:
        raise DegradeError(
            f"shuffled_numbers: {fmt} degraded output has an unchanged number multiset"
        )


# --------------------------------------------------------------------------
# Digit/number transforms shared by the shuffled_numbers builders.
# --------------------------------------------------------------------------


def _shuffle_single_digit(d: int) -> int:
    """Map a digit in 0..9 to a different digit in 1..9, deterministically."""
    return (d % 9) + 1


def _shuffle_positive_int(value: int) -> int:
    """Digit-rotate a non-negative int, guaranteed to change its value.

    Single-digit values (<=9) use the digit map above. Multi-digit values
    move their leading digit to the end (25307 -> 53072); an all-identical
    digit string (e.g. 111) would rotate to itself, so that case is handled
    by adding 1 instead.
    """
    if value <= 9:
        return _shuffle_single_digit(value)
    digits = str(value)
    if len(set(digits)) == 1:
        return int(digits) + 1
    return int(digits[1:] + digits[0])


def _split_float(value: float) -> tuple[int, str]:
    text = repr(value)
    if "e" in text or "E" in text:
        text = f"{value:.6f}"
    if "." in text:
        int_str, frac_str = text.split(".", 1)
        frac_str = frac_str.rstrip("0")
    else:
        int_str, frac_str = text, ""
    return int(int_str), frac_str


def _shuffle_number_value(value: int | float) -> int | float:
    """Transform a numeric xlsx cell value per the U7 shuffling rule."""
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        sign = -1 if value < 0 else 1
        return sign * _shuffle_positive_int(abs(value))
    if isinstance(value, float):
        sign = -1.0 if value < 0 else 1.0
        int_part, frac_str = _split_float(abs(value))
        new_int = _shuffle_positive_int(int_part)
        if frac_str:
            return sign * float(f"{new_int}.{frac_str}")
        return sign * float(new_int)
    return value


def _shuffle_number_token(token: str) -> str:
    """Transform a decimal-string number found in text (docx/pptx runs)."""
    if "." in token:
        int_str, frac_str = token.split(".", 1)
        new_int = _shuffle_positive_int(int(int_str))
        return f"{new_int}.{frac_str}"
    return str(_shuffle_positive_int(int(token)))


def _shuffle_text(text: str) -> str:
    return _NUMBER_TOKEN_RE.sub(lambda m: _shuffle_number_token(m.group(0)), text)


def _shuffle_paragraph_runs(paragraphs) -> None:
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.text and any(ch.isdigit() for ch in run.text):
                run.text = _shuffle_text(run.text)


# --------------------------------------------------------------------------
# roundtrip_control: identical load/save pipeline, no semantic change.
# --------------------------------------------------------------------------


def _control_xlsx(gold: Path, out_path: Path) -> Path:
    wb = openpyxl.load_workbook(gold)
    try:
        wb.save(out_path)
    finally:
        wb.close()
    return out_path


def _control_docx(gold: Path, out_path: Path) -> Path:
    doc = DocxDocument(str(gold))
    doc.save(out_path)
    return out_path


def _control_pptx(gold: Path, out_path: Path) -> Path:
    prs = Presentation(str(gold))
    prs.save(out_path)
    return out_path


def _control_pdf(gold: Path, out_path: Path) -> Path:
    reader = pypdf.PdfReader(str(gold))
    writer = pypdf.PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with out_path.open("wb") as fh:
        writer.write(fh)
    return out_path


# --------------------------------------------------------------------------
# truncated: drop the back half of the content, keep at least one unit.
# --------------------------------------------------------------------------


def _truncate_xlsx(gold: Path, out_path: Path) -> Path:
    # Drop the back half of the worksheets (real-data calibration showed
    # row-halving barely moves the judge on structure-heavy workbooks:
    # a whole missing sheet is what "half the work missing" looks like).
    # Single-sheet workbooks fall back to row-halving.
    wb = openpyxl.load_workbook(gold)
    try:
        n_sheets = len(wb.worksheets)
        if n_sheets > 1:
            drop, keep = _half_drop_keep(n_sheets)
            for ws in list(wb.worksheets[keep:]):
                wb.remove(ws)
            wb.save(out_path)
            _assert_count_reduced(n_sheets, keep, "xlsx", "sheet")
        else:
            ws = wb.worksheets[0]
            n = ws.max_row or 0
            drop, keep = _half_drop_keep(n)
            if drop > 0:
                ws.delete_rows(keep + 1, drop)
            wb.save(out_path)
            _assert_count_reduced(n, keep, "xlsx", "row")
    finally:
        wb.close()
    return out_path


def _truncate_docx(gold: Path, out_path: Path) -> Path:
    doc = DocxDocument(str(gold))
    body = doc.element.body
    block_tags = (docx_qn("w:p"), docx_qn("w:tbl"))
    children = [c for c in body if c.tag in block_tags]
    n = len(children)
    drop, keep = _half_drop_keep(n)
    for element in children[keep:]:
        body.remove(element)
    doc.save(out_path)
    _assert_count_reduced(n, keep, "docx", "paragraph")
    return out_path


def _truncate_pptx(gold: Path, out_path: Path) -> Path:
    prs = Presentation(str(gold))
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    n = len(slide_ids)
    drop, keep = _half_drop_keep(n)
    for sld_id in slide_ids[keep:]:
        r_id = sld_id.get(pptx_qn("r:id"))
        prs.part.drop_rel(r_id)
        slide_id_list.remove(sld_id)
    prs.save(out_path)
    _assert_count_reduced(n, keep, "pptx", "slide")
    return out_path


def _truncate_pdf(gold: Path, out_path: Path) -> Path:
    reader = pypdf.PdfReader(str(gold))
    n = len(reader.pages)
    drop, keep = _half_drop_keep(n)
    writer = pypdf.PdfWriter()
    for i in range(keep):
        writer.add_page(reader.pages[i])
    with out_path.open("wb") as fh:
        writer.write(fh)
    _assert_count_reduced(n, keep, "pdf", "page")
    return out_path


# --------------------------------------------------------------------------
# shuffled_numbers: change the numeric content, keep structure intact.
# pdf content streams can't be rewritten, so its variant instead drops the
# half of pages with the most digits (falling back to the truncated halving
# rule for a single-page pdf — which the shared self-check below will then
# correctly reject as unchanged, since a single page can't be reduced).
# --------------------------------------------------------------------------


def _shuffle_xlsx(gold: Path, out_path: Path) -> Path:
    wb = openpyxl.load_workbook(gold)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        continue
                    cell.value = _shuffle_number_value(value)
        wb.save(out_path)
    finally:
        wb.close()
    _assert_numbers_changed(gold, out_path, "xlsx")
    return out_path


def _shuffle_docx(gold: Path, out_path: Path) -> Path:
    doc = DocxDocument(str(gold))
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    _shuffle_paragraph_runs(paragraphs)
    doc.save(out_path)
    _assert_numbers_changed(gold, out_path, "docx")
    return out_path


def _shuffle_pptx(gold: Path, out_path: Path) -> Path:
    prs = Presentation(str(gold))
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                _shuffle_paragraph_runs(shape.text_frame.paragraphs)
    prs.save(out_path)
    _assert_numbers_changed(gold, out_path, "pptx")
    return out_path


def _shuffle_pdf(gold: Path, out_path: Path) -> Path:
    reader = pypdf.PdfReader(str(gold))
    pages = reader.pages
    n = len(pages)
    if n <= 1:
        _, keep = _half_drop_keep(n)
        keep_indices = list(range(keep))
    else:
        drop = n // 2
        digit_counts = [
            (len(re.findall(r"\d", page.extract_text() or "")), i)
            for i, page in enumerate(pages)
        ]
        ranked_by_digits = sorted(digit_counts, key=lambda t: t[0], reverse=True)
        drop_indices = {i for _, i in ranked_by_digits[:drop]}
        keep_indices = [i for i in range(n) if i not in drop_indices]

    writer = pypdf.PdfWriter()
    for i in keep_indices:
        writer.add_page(pages[i])
    with out_path.open("wb") as fh:
        writer.write(fh)
    _assert_numbers_changed(gold, out_path, "pdf")
    return out_path


_CONTROL_BUILDERS = {
    "xlsx": _control_xlsx,
    "docx": _control_docx,
    "pptx": _control_pptx,
    "pdf": _control_pdf,
}
_TRUNCATE_BUILDERS = {
    "xlsx": _truncate_xlsx,
    "docx": _truncate_docx,
    "pptx": _truncate_pptx,
    "pdf": _truncate_pdf,
}
_SHUFFLE_BUILDERS = {
    "xlsx": _shuffle_xlsx,
    "docx": _shuffle_docx,
    "pptx": _shuffle_pptx,
    "pdf": _shuffle_pdf,
}


def make_degraded(gold: Path, out_dir: Path) -> dict[str, Path]:
    """Produce the three calibration variants of `gold` inside `out_dir`.

    Returns a dict keyed by variant name ("roundtrip_control", "truncated",
    "shuffled_numbers") mapping to the written file's path. Every output
    keeps `gold`'s format and extension, named `{variant}_{gold.name}`.
    """
    fmt = gold.suffix.lower().lstrip(".")
    if fmt not in _SUPPORTED_FORMATS:
        raise DegradeError(f"unsupported deliverable extension: {fmt!r}")

    out_dir.mkdir(parents=True, exist_ok=True)

    return {
        "roundtrip_control": _CONTROL_BUILDERS[fmt](
            gold, out_dir / f"roundtrip_control_{gold.name}"
        ),
        "truncated": _TRUNCATE_BUILDERS[fmt](gold, out_dir / f"truncated_{gold.name}"),
        "shuffled_numbers": _SHUFFLE_BUILDERS[fmt](
            gold, out_dir / f"shuffled_numbers_{gold.name}"
        ),
    }
