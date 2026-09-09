"""使用合成任务和本地附件验收完整测评流程。"""

import csv
import json
from contextlib import contextmanager
from dataclasses import replace

import httpx
import pytest
from docx import Document

import evaluate as ev
import grade
import llm
from grade import Item, Task

PRODUCTS = ("豆包工作", "QwenWork", "WorkBuddy")


@pytest.fixture
def exam(tmp_path, monkeypatch):
    task = Task(
        1, "synthetic", "Occupation", "Synthetic task", (Item("a", "condition", 2),), (), ()
    )
    monkeypatch.setattr(ev, "DOWNLOADS", tmp_path / "downloads", raising=False)
    monkeypatch.setattr(ev, "RESULTS", tmp_path / "results", raising=False)
    return tmp_path / "tasks", [task]


def metrics(root, **updates):
    path = root / "metrics.csv"
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    rows[0].update(updates)
    grade.write_csv(path, rows)


def answer(root, text="answer", name="final.docx"):
    path = root / PRODUCTS[0] / "t_01" / "output" / name
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return path


def test_prepare_exports_only_inputs_and_preserves_manual_work(exam):
    root, tasks = exam
    ev.prepare(root, tasks)
    for product in PRODUCTS:
        packet = root / product / "t_01"
        assert (packet / "prompt.txt").read_text() == "Synthetic task"
        assert (packet / "output").is_dir()
        assert "condition" not in "".join(p.read_text() for p in packet.glob("*.txt"))
    saved = answer(root).read_bytes()
    metrics(root, status="completed", session="session-1", elapsed_seconds="0", tokens="0")
    ev.prepare(root, tasks)
    assert answer_path(root).read_bytes() == saved
    assert ev.read_metrics(root, tasks)[(PRODUCTS[0], 1)]["tokens"] == "0"


def answer_path(root):
    return root / PRODUCTS[0] / "t_01" / "output" / "final.docx"


def test_prepare_downloads_reference_once_for_three_identical_copies(exam, monkeypatch):
    root, tasks = exam
    task = replace(
        tasks[0],
        reference_names=("reference_files/data.txt",),
        reference_urls=("https://example.org/data.txt",),
    )
    calls = []

    def download(url, path):
        calls.append(url)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("source numbers")

    monkeypatch.setattr(ev, "download", download)
    ev.prepare(root, [task])
    ev.prepare(root, [task])
    assert calls == ["https://example.org/data.txt"]
    for product in PRODUCTS:
        assert (root / product / "t_01/reference_files/data.txt").read_text() == "source numbers"


@pytest.mark.parametrize(
    "name",
    [
        "../escape.txt",
        "/absolute.txt",
        "reference_files/../x.txt",
        "reference_files\\x.txt",
        "C:/x.txt",
    ],
)
def test_prepare_rejects_unsafe_reference_names(exam, name):
    root, tasks = exam
    task = replace(tasks[0], reference_names=(name,), reference_urls=("https://example.org/x",))
    with pytest.raises(ValueError):
        ev.prepare(root, [task])


def test_changed_exam_or_inputs_require_a_fresh_packet(exam):
    root, tasks = exam
    ev.prepare(root, tasks)
    with pytest.raises(ValueError):
        ev.prepare(root, [replace(tasks[0], task_id="another-task")])
    (root / PRODUCTS[0] / "t_01/prompt.txt").write_text("changed")
    with pytest.raises(ValueError):
        ev.prepare(root, tasks)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("elapsed_seconds", "-1"),
        ("elapsed_seconds", "nan"),
        ("elapsed_seconds", "inf"),
        ("elapsed_seconds", "bad"),
        ("tokens", "1.5"),
        ("tokens", "-1"),
        ("status", "unknown"),
        ("session", ""),
    ],
)
def test_invalid_manual_metrics_fail_before_judging(exam, field, value):
    root, tasks = exam
    ev.prepare(root, tasks)
    metrics(
        root,
        status="completed",
        session="session-1",
        **({field: value} if field not in {"status", "session"} else {}),
    )
    if field in {"status", "session"}:
        metrics(root, **{field: value})
    with pytest.raises(ValueError):
        ev.read_metrics(root, tasks)


@pytest.mark.parametrize("kind", ["duplicate", "missing", "wrong_product", "wrong_ordinal"])
def test_metrics_identity_is_complete_and_unique(exam, kind):
    root, tasks = exam
    ev.prepare(root, tasks)
    with (root / "metrics.csv").open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if kind == "duplicate":
        rows.append(rows[0])
    elif kind == "missing":
        rows.pop()
    elif kind == "wrong_product":
        rows[0]["product"] = "wrong"
    else:
        rows[0]["task"] = "0"
    grade.write_csv(root / "metrics.csv", rows)
    with pytest.raises(ValueError):
        ev.read_metrics(root, tasks)


def test_real_docx_scoring_cache_invalidation_and_pending_report(exam, monkeypatch):
    root, tasks = exam
    ev.prepare(root, tasks)
    answer(root, "first")
    metrics(root, status="completed", session="s1", elapsed_seconds="12.5", tokens="100")
    seen = []

    def ask(prompt, **kw):
        seen.append(prompt)
        return llm.Reply("MET", 0.01)

    monkeypatch.setattr(llm, "ask", ask)
    kwargs = {"model": "model-1", "modes": ("rubric",)}
    rows = ev.evaluate(root, tasks, **kwargs)
    assert rows[0]["rubric_score"] == 5.0
    assert rows[1]["rubric_score"] == ""
    assert rows[1]["status"] == "pending"
    ev.evaluate(root, tasks, **kwargs)
    assert len(seen) == 1
    answer(root, "second")
    ev.evaluate(root, tasks, **kwargs)
    assert len(seen) == 2
    ev.evaluate(root, tasks, model="model-2", modes=("rubric",))
    assert len(seen) == 3
    (root / PRODUCTS[0] / "t_01/prompt.txt").write_text("wrong task")
    rows = ev.evaluate(root, tasks, **kwargs)
    assert rows[0]["status"] == "error"
    assert rows[0]["rubric_score"] == ""
    assert len(seen) == 3


def test_reference_and_multiple_outputs_reach_judge(exam, monkeypatch):
    root, tasks = exam
    task = replace(
        tasks[0],
        reference_names=("reference_files/data.txt",),
        reference_urls=("https://example.org/data.txt",),
    )

    def download(url, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("reference evidence")

    monkeypatch.setattr(ev, "download", download)
    ev.prepare(root, [task])
    answer(root, "part one")
    answer(root, "part two", "other.docx")
    metrics(root, status="completed", session="s1")
    seen = []

    def ask(prompt, **kw):
        seen.append(prompt)
        return llm.Reply("MET", 0.0)

    monkeypatch.setattr(llm, "ask", ask)
    ev.evaluate(root, [task], model="m", modes=("rubric",))
    assert all(s in seen[0] for s in ["reference evidence", "part one", "part two"])


@pytest.mark.parametrize("problem", ["missing", "unsupported", "empty", "judge", "symlink"])
def test_infrastructure_errors_are_unscored(exam, monkeypatch, problem):
    root, tasks = exam
    ev.prepare(root, tasks)
    metrics(root, status="completed", session="s1")
    if problem == "unsupported":
        (root / PRODUCTS[0] / "t_01/output/x.zip").write_bytes(b"x")
    elif problem == "empty":
        answer(root, "")
    elif problem == "symlink":
        other = root / "other.docx"
        other.write_bytes(b"x")
        answer_path(root).symlink_to(other)
    elif problem == "judge":
        answer(root)

    def ask(*args, **kwargs):
        raise llm.JudgeError("Synthetic failure")

    monkeypatch.setattr(llm, "ask", ask)
    rows = ev.evaluate(root, tasks, model="m", modes=("rubric",))
    assert rows[0]["status"] == "error"
    assert rows[0]["rubric_score"] == ""


def test_failed_attempt_scores_zero_without_judge_and_missing_metrics_stay_blank(exam):
    root, tasks = exam
    ev.prepare(root, tasks)
    metrics(root, status="failed", session="s1", notes="timed out")
    rows = ev.evaluate(root, tasks, model="m", modes=("rubric", "pairwise"))
    assert rows[0]["rubric_score"] == 0
    assert rows[0]["pairwise_score"] == 0
    assert rows[0]["elapsed_seconds"] == ""
    assert rows[0]["tokens"] == ""


def test_report_ranks_only_complete_products_and_preserves_ties(tmp_path):
    rows = []
    for product, score in zip(PRODUCTS, [4, 4, ""], strict=True):
        for ordinal in [1, 2]:
            rows.append(
                dict(
                    product=product,
                    task=ordinal,
                    status="graded" if score else "pending",
                    rubric_score=score,
                    pairwise_score="",
                    elapsed_seconds="0",
                    tokens="",
                    session="s",
                    notes="",
                    judge_cost_usd=0.0,
                )
            )
    ev.write_reports(tmp_path, rows, 2, ("rubric",), "synthetic-model")
    with (tmp_path / "leaderboard.csv").open(encoding="utf-8-sig", newline="") as f:
        ranks = list(csv.DictReader(f))
    assert [r["rank"] for r in ranks] == ["1", "1", ""]
    assert ranks[0]["rubric_total"] == "8.0"
    assert ranks[0]["rubric_max"] == "10"
    assert ranks[0]["elapsed_seconds_total"] == "0.0"
    assert ranks[0]["tokens_total"] == ""
    assert ranks[2]["rubric_total"] == ""
    report = (tmp_path / "report.md").read_text()
    assert "synthetic-model" in report
    assert "逐题" in report


def test_manifest_contains_only_expected_input_paths(exam):
    root, tasks = exam
    ev.prepare(root, tasks)
    manifest = json.loads((root / "manifest.json").read_text())
    assert len(manifest["tasks"]) == 1
    assert "prompt.txt" in manifest["tasks"][0]["files"]
    assert all("output" not in p for p in manifest["tasks"][0]["files"])


def test_reusing_one_session_for_two_questions_is_rejected(exam):
    root, tasks = exam
    tasks.append(replace(tasks[0], ordinal=2, task_id="synthetic-2"))
    ev.prepare(root, tasks)
    with (root / "metrics.csv").open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    for row in rows[:2]:
        row.update(status="completed", session="same-session")
    grade.write_csv(root / "metrics.csv", rows)
    with pytest.raises(ValueError):
        ev.read_metrics(root, tasks)


def test_both_modes_use_correct_gold_and_resume_without_network_judge(exam, monkeypatch):
    root, tasks = exam
    ev.prepare(root, tasks)
    answer(root, "product evidence")
    metrics(root, status="completed", session="s1")
    monkeypatch.setattr(grade, "gold_text", lambda t: f"expert for {t.ordinal}")
    calls = []

    def ask(prompt, **kw):
        calls.append(prompt)
        return llm.Reply("MET" if kw["system"] == grade.RUBRIC_SYSTEM else "TIE", 0.02)

    monkeypatch.setattr(llm, "ask", ask)
    rows = ev.evaluate(root, tasks, model="m", modes=("rubric", "pairwise"))
    assert rows[0]["rubric_score"] == 5
    assert rows[0]["pairwise_score"] == 0.5
    assert rows[0]["judge_cost_usd"] == pytest.approx(0.04)
    assert "expert for 1" in calls[1]
    ev.evaluate(root, tasks, model="m", modes=("rubric", "pairwise"))
    assert len(calls) == 2


def test_limit_keeps_full_exam_denominator(exam):
    root, tasks = exam
    tasks.append(replace(tasks[0], ordinal=2, task_id="synthetic-2"))
    ev.prepare(root, tasks)
    metrics(root, status="failed", session="s1")
    rows = ev.evaluate(root, tasks, model="m", modes=("rubric",), limit=1)
    assert rows[1]["status"] == "not_evaluated"
    with (ev.report_dir(root) / "leaderboard.csv").open(encoding="utf-8-sig", newline="") as f:
        ranks = list(csv.DictReader(f))
    assert ranks[0]["total_tasks"] == "2"
    assert ranks[0]["rubric_total"] == ""
    assert ranks[0]["rank"] == ""


@pytest.mark.parametrize("value", ["1e309", "9" * 5000], ids=["exponent", "oversized-integer"])
def test_extreme_manual_numbers_fail_preflight(exam, value):
    root, tasks = exam
    ev.prepare(root, tasks)
    metrics(root, tokens=value)
    with pytest.raises(ValueError):
        ev.read_metrics(root, tasks)


def test_download_retries_transient_error_and_publishes_complete_file(tmp_path, monkeypatch):
    calls = []

    @contextmanager
    def stream(*args, **kwargs):
        calls.append(1)
        yield httpx.Response(
            503 if len(calls) == 1 else 200,
            content=b"complete",
            request=httpx.Request("GET", "https://example.org/source"),
        )

    monkeypatch.setattr(httpx, "stream", stream)
    monkeypatch.setattr(ev, "sleep", lambda _: None, raising=False)
    path = tmp_path / "download.txt"
    ev.download("https://example.org/source", path)
    assert path.read_bytes() == b"complete"
    assert len(calls) == 2
    assert not path.with_suffix(".txt.part").exists()


def test_desktop_metadata_and_office_locks_do_not_become_answers(exam, monkeypatch):
    root, tasks = exam
    ev.prepare(root, tasks)
    answer(root)
    metrics(root, status="completed", session="s1")
    packet = root / PRODUCTS[0] / "t_01"
    (packet / "reference_files/.DS_Store").write_bytes(b"metadata")
    (packet / "output/.DS_Store").write_bytes(b"metadata")
    (packet / "output/~$final.docx").write_bytes(b"lock")
    monkeypatch.setattr(llm, "ask", lambda *a, **k: llm.Reply("MET", 0))
    rows = ev.evaluate(root, tasks, model="m", modes=("rubric",))
    assert rows[0]["status"] == "graded"
    assert rows[0]["rubric_score"] == 5


def test_cli_full_three_product_flow_uses_real_llm_client_with_mock_http(exam, monkeypatch):
    root, tasks = exam
    monkeypatch.setattr(grade, "load_tasks", lambda: tasks)
    monkeypatch.setattr(grade, "gold_text", lambda t: "expert evidence")
    monkeypatch.setenv("OPENROUTER_API_KEY", "synthetic-key")
    assert ev.main(["prepare", "--root", str(root)]) == 0
    with (root / "metrics.csv").open(encoding="utf-8-sig", newline="") as f:
        records = list(csv.DictReader(f))
    for i, product in enumerate(PRODUCTS):
        doc = Document()
        doc.add_paragraph("strong answer" if i != 1 else "weak answer")
        doc.save(root / product / "t_01/output/final.docx")
        records[i].update(
            status="completed", session=f"session-{i}", elapsed_seconds="10", tokens="20"
        )
    grade.write_csv(root / "metrics.csv", records)
    requests = []

    def handler(request):
        body = json.loads(request.content)
        requests.append(body)
        assert body["provider"] == llm.PROVIDER
        prompt = body["messages"][-1]["content"]
        is_rubric = body["messages"][0]["content"] == grade.RUBRIC_SYSTEM
        verdict = ("MET" if "strong answer" in prompt else "NOT_MET") if is_rubric else "TIE"
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": verdict}}],
                "provider": "DeepSeek",
                "usage": {"cost": 0.001},
            },
        )

    client = httpx.Client
    monkeypatch.setattr(httpx, "Client", lambda: client(transport=httpx.MockTransport(handler)))
    assert ev.main(["run", "--root", str(root)]) == 0
    assert len(requests) == 6
    with (ev.report_dir(root) / "leaderboard.csv").open(encoding="utf-8-sig", newline="") as f:
        ranks = list(csv.DictReader(f))
    assert [(r["product"], r["rank"], r["rubric_total"]) for r in ranks] == [
        (PRODUCTS[0], "1", "5.0"),
        (PRODUCTS[2], "1", "5.0"),
        (PRODUCTS[1], "3", "0.0"),
    ]
    assert all(r["pairwise_win_rate"] == "0.5" for r in ranks)
    assert all(r["tokens_total"] == "20" for r in ranks)
    assert ev.main(["run", "--root", str(root)]) == 0
    assert len(requests) == 6


def hashed_reference(exam, monkeypatch):
    root, tasks = exam
    task = replace(
        tasks[0],
        reference_names=(f"reference_files/{'a' * 32}/source.txt",),
        reference_urls=("https://example.org/source.txt",),
    )

    def download(url, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("source evidence")

    monkeypatch.setattr(ev, "download", download)
    return root, [task]


def test_reference_export_uses_original_filename_directly(exam, monkeypatch):
    root, tasks = hashed_reference(exam, monkeypatch)
    ev.prepare(root, tasks)
    for product in PRODUCTS:
        reference = root / product / "t_01/reference_files"
        assert sorted(p.name for p in reference.iterdir()) == ["source.txt"]
        assert (reference / "source.txt").read_text() == "source evidence"
    record = ev.read_manifest(root, tasks)["tasks"][0]
    assert "reference_files/source.txt" in record["files"]


def test_existing_hash_layout_is_flattened_with_answers_and_metrics_preserved(exam, monkeypatch):
    root, tasks = hashed_reference(exam, monkeypatch)
    with monkeypatch.context() as old:
        old.setattr(ev, "reference_path", lambda name: name)
        ev.prepare(root, tasks)
    original_answer = answer(root).read_bytes()
    metrics(root, status="completed", session="session-1", elapsed_seconds="9", tokens="17")
    original_metrics = (root / "metrics.csv").read_bytes()
    original_prompt = (root / PRODUCTS[0] / "t_01/prompt.txt").read_bytes()

    def no_download(*args):
        pytest.fail("Cached reference should be reused")

    monkeypatch.setattr(ev, "download", no_download)
    ev.prepare(root, tasks)
    ev.prepare(root, tasks)
    assert (root / "metrics.csv").read_bytes() == original_metrics
    assert answer_path(root).read_bytes() == original_answer
    assert (root / PRODUCTS[0] / "t_01/prompt.txt").read_bytes() == original_prompt
    record = ev.read_manifest(root, tasks)["tasks"][0]
    for product in PRODUCTS:
        packet = root / product / "t_01"
        assert not (packet / "reference_files" / ("a" * 32)).exists()
        ev.check_inputs(packet, record)


@pytest.mark.parametrize("second_name", ["source.txt", "SOURCE.txt"])
def test_flattened_filename_collision_is_reported_before_export(exam, monkeypatch, second_name):
    root, tasks = hashed_reference(exam, monkeypatch)
    task = replace(
        tasks[0],
        reference_names=(*tasks[0].reference_names, f"reference_files/{'b' * 32}/{second_name}"),
        reference_urls=(*tasks[0].reference_urls, "https://example.org/second.txt"),
    )
    with pytest.raises(ValueError, match="同名"):
        ev.prepare(root, [task])
    assert not (root / "manifest.json").exists()


@pytest.mark.parametrize("changed", ["legacy", "flat"])
def test_flattening_preserves_conflicting_local_files(exam, monkeypatch, changed):
    root, tasks = hashed_reference(exam, monkeypatch)
    with monkeypatch.context() as old:
        old.setattr(ev, "reference_path", lambda name: name)
        ev.prepare(root, tasks)
    packet = root / PRODUCTS[0] / "t_01"
    path = packet / (
        tasks[0].reference_names[0] if changed == "legacy" else "reference_files/source.txt"
    )
    path.write_text("human edit")
    manifest_before = (root / "manifest.json").read_bytes()
    with pytest.raises(ValueError):
        ev.prepare(root, tasks)
    assert path.read_text() == "human edit"
    assert (root / "manifest.json").read_bytes() == manifest_before
