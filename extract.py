"""把办公文件转成判分用的纯文本。支持 xlsx / docx / pptx / pdf。

异常消息只带格式名和数字,不带文件内容。
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pypdf
from defusedxml.ElementTree import fromstring
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

MAX_CHARS = 2_000_000
_MAX_CELLS = 200_000
_MAX_PDF_PAGES = 500


class ExtractError(Exception):
    pass


def extract(path: Path, *, max_chars: int = MAX_CHARS) -> str:
    """读一份交付物,返回纯文本(超长截断到 max_chars)。"""
    fmt = Path(path).suffix.lower().lstrip(".")
    readers = {"xlsx": _xlsx, "docx": _docx, "pptx": _pptx, "pdf": _pdf}
    if fmt not in readers:
        raise ExtractError(f"不支持的格式: {fmt!r}")
    try:
        text = _pdf(Path(path), max_chars) if fmt == "pdf" else readers[fmt](Path(path))
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError(f"{fmt}: 解析失败 ({type(exc).__name__})") from exc
    return text[:max_chars]


def _xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        lines = [f"[工作表] {', '.join(wb.sheetnames)}"]
        cells = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    cells += 1
                    if cells > _MAX_CELLS:
                        raise ExtractError(f"xlsx: 单元格数超过上限 {_MAX_CELLS}")
                    if cell.value is not None:
                        lines.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
        return "\n".join(lines)
    finally:
        wb.close()


def _docx(path: Path) -> str:
    # 正文和表格来自 document.xml；读取文本无需加载媒体与外部链接关系。
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as package:
        body = fromstring(package.read("word/document.xml")).find("w:body", ns)
    if body is None:
        raise ExtractError("docx: 缺少正文")

    def paragraph(node):
        parts = []
        for n in node.iter():
            if n.tag == f"{{{ns['w']}}}t":
                parts.append(n.text or "")
            elif n.tag == f"{{{ns['w']}}}tab":
                parts.append("\t")
            elif n.tag in {f"{{{ns['w']}}}br", f"{{{ns['w']}}}cr"}:
                parts.append("\n")
        return "".join(parts)

    lines = []
    for node in body:
        if node.tag == f"{{{ns['w']}}}p":
            text = paragraph(node)
            if text:
                lines.append(text)
        elif node.tag == f"{{{ns['w']}}}tbl":
            for row in node.findall("w:tr", ns):
                cells = [
                    "\n".join(paragraph(p) for p in cell.findall(".//w:p", ns))
                    for cell in row.findall("w:tc", ns)
                ]
                lines.append("[表格] " + " | ".join(cells))
    return "\n".join(lines)


def _pptx(path: Path) -> str:
    prs = Presentation(str(path))
    lines: list[str] = []

    def walk(shapes, n: int) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, n)
                continue
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = " | ".join(c.text for c in row.cells)
                    if row_text.strip():
                        lines.append(f"第{n}页[表格] {row_text}")
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text:
                lines.append(f"第{n}页: {shape.text_frame.text}")
            if getattr(shape, "has_chart", False):
                lines.append(f"第{n}页[含图表]")

    for n, slide in enumerate(prs.slides, start=1):
        walk(slide.shapes, n)
    return "\n".join(lines)


def _pdf(path: Path, max_chars: int = MAX_CHARS) -> str:
    reader = pypdf.PdfReader(str(path))
    if len(reader.pages) > _MAX_PDF_PAGES:
        raise ExtractError(f"pdf: 页数 {len(reader.pages)} 超过上限 {_MAX_PDF_PAGES}")
    lines, total = [], 0
    for page in reader.pages:
        t = page.extract_text() or ""
        if t:
            lines.append(t)
            total += len(t)
        if total > max_chars:
            break
    return "\n".join(lines)
