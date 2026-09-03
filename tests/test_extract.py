"""Tests for extract_document (spec U1).

All fixtures below are tiny synthetic documents generated on the fly with
openpyxl/python-docx/python-pptx/pypdf; none originate from the real
GDPval dataset (see AGENTS.md secrecy rules). No file under secrets/ or
data/ is read.
"""

from __future__ import annotations

import struct
import zipfile
import zlib
from pathlib import Path

import openpyxl
import pytest
from docx import Document as DocxDocument
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from gdpval_eval import extract
from gdpval_eval.extract import ExtractError, extract_document


def _tiny_png_bytes() -> bytes:
    """A syntactically valid 1x1 red PNG, for embedding as an inline image."""
    width, height = 1, 1

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\x00\x00")
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def _add_pdf_text_page(writer: PdfWriter, text: str) -> None:
    """Add a page with a real, extractable text content stream (no reportlab)."""
    page = writer.add_blank_page(width=300, height=300)

    font_dict = DictionaryObject()
    font_dict[NameObject("/Type")] = NameObject("/Font")
    font_dict[NameObject("/Subtype")] = NameObject("/Type1")
    font_dict[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font_dict)

    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_ref
    resources = DictionaryObject()
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources

    stream_obj = DecodedStreamObject()
    stream_obj.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode("latin-1"))
    page.replace_contents(stream_obj)


def _build_entity_bomb_xlsx(path: Path) -> None:
    """A valid-looking xlsx whose one worksheet part hides an XML entity bomb.

    Built by taking a real openpyxl-produced xlsx apart and swapping in a
    hand-written sheet XML with a DOCTYPE/ENTITY declaration, so the zip
    safety pre-check (entry count/size) passes and only the OOXML XML
    parser sees the bomb.
    """
    seed_path = path.with_suffix(".seed.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "seed"
    wb.save(seed_path)

    malicious_sheet = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE worksheet [
<!ENTITY a "AAAAAAAAAA">
<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">
<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">
]>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&c;</t></is></c></row></sheetData>
</worksheet>
"""

    with zipfile.ZipFile(seed_path) as src, zipfile.ZipFile(path, "w") as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = malicious_sheet
            dst.writestr(name, data)

    seed_path.unlink()


def test_extract_xlsx_parses_values_formulas_and_chart(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws["A1"] = 10
    ws["A2"] = 32
    ws["B2"] = "=A1+A2"
    chart = BarChart()
    data = Reference(ws, min_col=1, min_row=1, max_row=2)
    chart.add_data(data)
    ws.add_chart(chart, "D5")
    wb.save(path)

    doc = extract_document(path)

    assert doc.fmt == "xlsx"
    assert doc.facts.sheet_names == ("Ledger",)
    assert doc.facts.page_count is None
    assert doc.facts.slide_count is None
    assert doc.facts.has_chart is True
    assert doc.facts.formula_count == 1
    assert doc.facts.truncated is False
    assert doc.facts.text_chars == len(doc.text)
    assert "Ledger!A1: 10" in doc.text
    assert "Ledger!A2: 32" in doc.text


def test_extract_docx_parses_paragraphs_tables_and_images(tmp_path):
    path = tmp_path / "memo.docx"
    png_path = tmp_path / "tiny.png"
    png_path.write_bytes(_tiny_png_bytes())

    doc = DocxDocument()
    doc.add_paragraph("Synthetic paragraph about a fictitious project.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "left-cell"
    table.rows[0].cells[1].text = "right-cell"
    doc.add_picture(str(png_path))
    doc.save(path)

    extracted = extract_document(path)

    assert extracted.fmt == "docx"
    assert extracted.facts.page_count is None
    assert extracted.facts.sheet_names == ()
    assert extracted.facts.image_count == 1
    assert extracted.facts.text_chars == len(extracted.text)
    assert "Synthetic paragraph about a fictitious project." in extracted.text
    assert "left-cell" in extracted.text
    assert "right-cell" in extracted.text


def test_extract_pptx_parses_slides_and_counts(tmp_path):
    path = tmp_path / "deck.pptx"
    png_path = tmp_path / "tiny.png"
    png_path.write_bytes(_tiny_png_bytes())

    prs = Presentation()
    layout = prs.slide_layouts[5]

    slide1 = prs.slides.add_slide(layout)
    slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (
        "First fictitious slide body."
    )

    slide2 = prs.slides.add_slide(layout)
    slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (
        "Second fictitious slide body."
    )
    chart_data = CategoryChartData()
    chart_data.categories = ["x", "y"]
    chart_data.add_series("series-a", (1, 2))
    slide2.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(3), Inches(2), chart_data
    )
    slide2.shapes.add_picture(str(png_path), Inches(1), Inches(4), Inches(0.2), Inches(0.2))

    prs.save(path)

    extracted = extract_document(path)

    assert extracted.fmt == "pptx"
    assert extracted.facts.slide_count == 2
    assert extracted.facts.page_count is None
    assert extracted.facts.has_chart is True
    assert extracted.facts.image_count == 1
    assert extracted.facts.text_chars == len(extracted.text)
    assert "First fictitious slide body." in extracted.text
    assert "Second fictitious slide body." in extracted.text


def test_extract_pdf_pages_and_zero_page_text(tmp_path):
    real_path = tmp_path / "report.pdf"
    writer = PdfWriter()
    _add_pdf_text_page(writer, "First synthetic page body")
    _add_pdf_text_page(writer, "Second synthetic page body")
    with real_path.open("wb") as fh:
        writer.write(fh)

    real_doc = extract_document(real_path)

    assert real_doc.fmt == "pdf"
    assert real_doc.facts.page_count == 2
    assert real_doc.facts.sheet_names == ()
    assert real_doc.facts.slide_count is None
    assert "First synthetic page body" in real_doc.text
    assert "Second synthetic page body" in real_doc.text
    assert real_doc.facts.text_chars == len(real_doc.text)

    empty_path = tmp_path / "empty.pdf"
    empty_writer = PdfWriter()
    with empty_path.open("wb") as fh:
        empty_writer.write(fh)

    empty_doc = extract_document(empty_path)

    assert empty_doc.fmt == "pdf"
    assert empty_doc.facts.page_count == 0
    assert empty_doc.text == ""
    assert empty_doc.facts.text_chars == 0
    assert empty_doc.facts.truncated is False


def test_extract_xlsx_entity_bomb_zip_rejected(tmp_path):
    path = tmp_path / "bomb.xlsx"
    _build_entity_bomb_xlsx(path)

    with pytest.raises(ExtractError):
        extract_document(path)


def test_extract_xlsx_zip_total_size_exceeds_limit_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "_MAX_ZIP_TOTAL_BYTES", 32)

    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "small workbook, far above 32 bytes uncompressed"
    wb.save(path)

    with pytest.raises(ExtractError):
        extract_document(path)


def test_extract_xlsx_cell_cap_truncates_without_raising(tmp_path, monkeypatch):
    monkeypatch.setattr(extract, "_MAX_XLSX_CELLS", 2)

    path = tmp_path / "grid.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in range(1, 4):
        for col in range(1, 3):
            ws.cell(row=row, column=col, value=row * 10 + col)
    wb.save(path)

    doc = extract_document(path)

    assert doc.facts.truncated is True
    assert len(doc.text.splitlines()) <= 2


def test_extract_unsupported_extension_raises(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("plain text, not a supported deliverable format")

    with pytest.raises(ExtractError):
        extract_document(path)
