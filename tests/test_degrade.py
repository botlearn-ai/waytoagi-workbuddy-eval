"""Tests for make_degraded, the calibration-gauge degrade tool (spec U7).

All fixtures are tiny synthetic documents generated on the fly with
openpyxl/python-docx/python-pptx/pypdf; none originate from the real
GDPval dataset (see AGENTS.md secrecy rules). No file under secrets/ or
data/ is read.

U7 is the "gauge" behind the discrimination check: every non-control
variant must be internally verified to differ from the gold file, so
several tests here specifically exercise the cases where a degraded
variant could otherwise end up accidentally equivalent to the original.
"""

from __future__ import annotations

import re

import openpyxl
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from gdpval_eval.degrade import DegradeError, _shuffle_pdf, make_degraded
from gdpval_eval.extract import extract_document

_NUMBER_RE = re.compile(r"\d+(?:\.\d+)?")


def _add_pdf_text_page(writer: PdfWriter, text: str) -> None:
    """Add a page with a real, extractable text content stream (no reportlab)."""
    page = writer.add_blank_page(width=300, height=300)

    font_dict = DictionaryObject()
    font_dict[NameObject("/Type")] = NameObject("/Font")
    font_dict[NameObject("/Subtype")] = NameObject("/Type1")
    font_dict[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font_dict)

    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_ref
    resources = DictionaryObject()
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources

    stream_obj = DecodedStreamObject()
    stream_obj.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode("latin-1"))
    page.replace_contents(stream_obj)


def test_degrade_xlsx_truncated_halves_rows_and_control_preserves_values(tmp_path):
    gold = tmp_path / "ledger.xlsx"
    wb = openpyxl.Workbook()
    ledger = wb.active
    ledger.title = "Ledger"
    for i, value in enumerate([111, 25307, 481, 9], start=1):
        ledger.cell(row=i, column=1, value=value)
    tiny = wb.create_sheet("Tiny")
    tiny["A1"] = 7
    wb.save(gold)

    result = make_degraded(gold, tmp_path / "out")

    assert set(result) == {"roundtrip_control", "truncated", "shuffled_numbers"}
    assert result["truncated"].name == f"truncated_{gold.name}"
    assert result["roundtrip_control"].name == f"roundtrip_control_{gold.name}"

    truncated_wb = openpyxl.load_workbook(result["truncated"])
    assert truncated_wb["Ledger"].max_row == 2
    assert [c.value for c in truncated_wb["Ledger"]["A"] if c.value is not None] == [111, 25307]
    assert truncated_wb["Tiny"].max_row == 1
    assert truncated_wb["Tiny"]["A1"].value == 7
    assert extract_document(result["truncated"]).fmt == "xlsx"

    control_wb = openpyxl.load_workbook(result["roundtrip_control"])
    assert [c.value for c in control_wb["Ledger"]["A"] if c.value is not None] == [
        111,
        25307,
        481,
        9,
    ]
    assert control_wb["Tiny"]["A1"].value == 7
    assert extract_document(result["roundtrip_control"]).fmt == "xlsx"


def test_degrade_xlsx_shuffled_numbers_changes_values_and_parses(tmp_path):
    gold = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 111
    ws["A2"] = 25307
    ws["A3"] = 9
    wb.save(gold)

    result = make_degraded(gold, tmp_path / "out")

    shuffled_wb = openpyxl.load_workbook(result["shuffled_numbers"])
    shuffled_ws = shuffled_wb["Data"]
    assert shuffled_ws["A1"].value == 112
    assert shuffled_ws["A2"].value == 53072
    assert shuffled_ws["A3"].value == 1

    gold_numbers = _NUMBER_RE.findall(extract_document(gold).text)
    out_numbers = _NUMBER_RE.findall(extract_document(result["shuffled_numbers"]).text)
    assert sorted(gold_numbers) != sorted(out_numbers)
    assert extract_document(result["shuffled_numbers"]).fmt == "xlsx"


def test_degrade_pptx_truncated_halves_slide_count(tmp_path):
    gold = tmp_path / "deck.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[5]
    for text in (
        "First slide body 12.",
        "Second slide body 34.",
        "Third slide body 56.",
        "Fourth slide body 78.",
    ):
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = text
    prs.save(gold)

    result = make_degraded(gold, tmp_path / "out")

    gold_doc = extract_document(gold)
    truncated_doc = extract_document(result["truncated"])
    assert gold_doc.facts.slide_count == 4
    assert truncated_doc.facts.slide_count == 2
    assert "First slide body" in truncated_doc.text
    assert "Second slide body" in truncated_doc.text
    assert "Third slide body" not in truncated_doc.text
    assert "Fourth slide body" not in truncated_doc.text


def test_degrade_pdf_truncated_halves_page_count(tmp_path):
    gold = tmp_path / "report.pdf"
    writer = PdfWriter()
    for text in (
        "Page one body 12.",
        "Page two body 34.",
        "Page three body 56.",
        "Page four body 78.",
    ):
        _add_pdf_text_page(writer, text)
    with gold.open("wb") as fh:
        writer.write(fh)

    result = make_degraded(gold, tmp_path / "out")

    gold_doc = extract_document(gold)
    truncated_doc = extract_document(result["truncated"])
    assert gold_doc.facts.page_count == 4
    assert truncated_doc.facts.page_count == 2
    assert "Page one body" in truncated_doc.text
    assert "Page two body" in truncated_doc.text
    assert "Page three body" not in truncated_doc.text


def test_degrade_pdf_single_page_shuffled_raises_degrade_error(tmp_path):
    gold = tmp_path / "single.pdf"
    writer = PdfWriter()
    _add_pdf_text_page(writer, "Report total 4821 units shipped this quarter.")
    with gold.open("wb") as fh:
        writer.write(fh)

    with pytest.raises(DegradeError):
        _shuffle_pdf(gold, tmp_path / "shuffled_numbers_single.pdf")


def test_degrade_docx_truncated_halves_body_elements(tmp_path):
    gold = tmp_path / "memo.docx"
    doc = DocxDocument()
    doc.add_paragraph("Paragraph one.")
    doc.add_paragraph("Paragraph two.")
    doc.add_paragraph("Paragraph three.")
    doc.add_paragraph("Paragraph four covers 42 items.")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "table cell"
    doc.save(gold)

    result = make_degraded(gold, tmp_path / "out")

    truncated = DocxDocument(str(result["truncated"]))
    assert [p.text for p in truncated.paragraphs] == [
        "Paragraph one.",
        "Paragraph two.",
        "Paragraph three.",
    ]
    assert len(truncated.tables) == 0
    assert extract_document(result["truncated"]).fmt == "docx"


def test_degrade_docx_shuffled_numbers_changes_text(tmp_path):
    gold = tmp_path / "figures.docx"
    doc = DocxDocument()
    doc.add_paragraph(
        "Revenue rose from 25307 to 481, a gain of 9 percent, across 111 branches."
    )
    doc.add_paragraph("End of report.")
    doc.save(gold)

    result = make_degraded(gold, tmp_path / "out")

    shuffled_text = extract_document(result["shuffled_numbers"]).text
    assert shuffled_text == (
        "Revenue rose from 53072 to 814, a gain of 1 percent, across 112 branches.\n"
        "End of report."
    )
    gold_numbers = _NUMBER_RE.findall(extract_document(gold).text)
    out_numbers = _NUMBER_RE.findall(shuffled_text)
    assert sorted(gold_numbers) != sorted(out_numbers)


def test_degrade_shuffled_no_digits_raises_degrade_error(tmp_path):
    gold = tmp_path / "prose.docx"
    doc = DocxDocument()
    doc.add_paragraph("Synthetic memo about a fictitious plan.")
    doc.add_paragraph("It contains no figures whatsoever.")
    doc.add_paragraph("Every sentence avoids numerals entirely.")
    doc.save(gold)

    with pytest.raises(DegradeError):
        make_degraded(gold, tmp_path / "out")
